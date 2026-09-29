"""
KiotViet Auto Export Bot - Danh sách vận đơn

Mỗi ngày bot tự:
  1. Đăng nhập KiotViet, vào trang Vận đơn, chọn đúng các cột cần lấy
  2. Chia tháng hiện tại thành các đợt 10 ngày (KiotViet giới hạn dung lượng mỗi lần xuất)
  3. Xuất từng đợt, chờ tải xong, gộp thành 1 file Excel (giữ nguyên định dạng Table)
  4. Chuyển file vào thư mục đích (ví dụ thư mục OneDrive mà Power BI đọc)
  5. Gửi email báo thành công hoặc báo lỗi kèm toàn bộ log

Tác giả : Thanh Ha
Phiên bản: 2.7
"""

import glob
import html
import os
import re
import shutil
import smtplib
import time
import traceback
from datetime import datetime, timedelta
from email.mime.multipart import MIMEMultipart
from email.mime.text import MIMEText

from dotenv import load_dotenv
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from selenium import webdriver
from selenium.webdriver.chrome.service import Service
from selenium.webdriver.common.by import By
from selenium.webdriver.common.keys import Keys
from selenium.webdriver.support import expected_conditions as EC
from selenium.webdriver.support.ui import WebDriverWait
from webdriver_manager.chrome import ChromeDriverManager

# =====================================================
# CẤU HÌNH (đọc từ file .env, xem .env.example)
# =====================================================
load_dotenv()

JOB_NAME = "KIOTVIET_VANDON"
LOGIN_URL = os.getenv("KIOTVIET_LOGIN_URL")          # https://<ten-cua-hang>.kiotviet.vn/man/#/login
USERNAME = os.getenv("KIOTVIET_USERNAME")
PASSWORD = os.getenv("KIOTVIET_PASSWORD")
DOWNLOAD_DIR = os.getenv("DOWNLOAD_DIR", os.path.join(os.path.expanduser("~"), "Downloads"))
FINAL_DIR = os.getenv("FINAL_DIR")                    # thư mục nhận file cuối
FILE_PREFIX = os.getenv("FILE_PREFIX", "DanhSachVanDon")
EMAIL_FROM = os.getenv("EMAIL_FROM")
EMAIL_PASSWORD = os.getenv("EMAIL_PASSWORD")          # Gmail App Password
EMAIL_TO = os.getenv("EMAIL_TO") or EMAIL_FROM

REQUIRED_COLUMNS = [
    "Mã vận đơn", "Thời gian tạo", "Thời gian hoàn thành",
    "Mã hóa đơn", "Khách hàng", "Điện thoại",
    "Địa chỉ", "Khu vực", "Trạng thái giao",
]
MAX_MISSING_COLUMNS = 1        # cho phép thiếu tối đa 1 cột (KiotViet đôi khi đổi tên cột)
CHUNK_DAYS = 10                # mỗi đợt xuất 10 ngày (KiotViet giới hạn dung lượng mỗi lần xuất)
DATA_WAIT_SECONDS = 600        # chờ lưới dữ liệu hiện ra tối đa 10 phút
EXPORT_WAIT_SECONDS = 3600     # chờ KiotViet tạo file xuất tối đa 1 giờ
DOWNLOAD_WAIT_SECONDS = 600    # chờ Chrome tải file tối đa 10 phút

LOG_BUFFER = []


def check_config():
    missing = [k for k, v in {
        "KIOTVIET_LOGIN_URL": LOGIN_URL, "KIOTVIET_USERNAME": USERNAME, "KIOTVIET_PASSWORD": PASSWORD,
        "FINAL_DIR": FINAL_DIR, "EMAIL_FROM": EMAIL_FROM, "EMAIL_PASSWORD": EMAIL_PASSWORD,
    }.items() if not v]
    if missing:
        raise ValueError(f"Thiếu cấu hình trong .env: {', '.join(missing)}")


# =====================================================
# LOG
# =====================================================
STYLES = {
    "header":  lambda m: f"\n{'═' * len(m)}\n{m}\n{'═' * len(m)}",
    "section": lambda m: f"\n{m}\n{'─' * len(m)}",
    "step":    lambda m: m,
    "info":    lambda m: m,
    "bullet":  lambda m: f"   • {m}",
    "loading": lambda m: f"⏳ {m}",
    "success": lambda m: f"✅ {m}",
    "warning": lambda m: f"⚠️  {m}",
    "result":  lambda m: f"\n{m}",
}


def log(msg, style="info"):
    """In ra màn hình và lưu vào bộ đệm để gửi kèm email."""
    text = STYLES[style](msg)
    print(text)
    LOG_BUFFER.append(text)


# =====================================================
# EMAIL
# =====================================================
EMAIL_TEMPLATE = """
<html><head><style>
  body {{ font-family: 'Courier New', monospace; margin: 0; background: #fff; }}
  .header {{ background: {color}; color: #fff; padding: 30px; text-align: center; }}
  .header h1 {{ margin: 0; font-size: 30px; }}
  .wrap {{ background: #f5f7fa; padding: 10px; }}
  .card {{ background: #fff; padding: 30px; margin: 0 auto 10px; max-width: 840px; border-radius: 12px;
           box-shadow: 0 2px 8px rgba(0,0,0,.15); }}
  .card h3 {{ margin-top: 0; border-bottom: 2px solid {color}; padding-bottom: 10px; }}
  .item {{ margin-bottom: 12px; line-height: 1.6; }}
  .label {{ font-weight: bold; color: #555; }}
  .log {{ background: #1e1e1e; color: #d4d4d4; padding: 30px; margin: 0 auto; max-width: 840px; border-radius: 12px; }}
  .log pre {{ background: #0d0d0d; padding: 20px; border-radius: 4px; overflow-x: auto; font-size: 13px; }}
  .footer {{ background: #263238; color: #aaa; padding: 20px; text-align: center; font-size: 12px; }}
</style></head><body>
  <div class="header"><h1>{title}</h1><h2>{job}</h2></div>
  <div class="wrap">
    <div class="card"><h3>TỔNG KẾT</h3>{items}</div>
    <div class="log"><h3>CHI TIẾT THỰC THI</h3><pre>{log}</pre></div>
  </div>
  <div class="footer">Email tự động từ {job}</div>
</body></html>
"""


def send_email(ok, summary):
    """Gửi email báo kết quả (xanh = thành công, đỏ = lỗi) kèm toàn bộ log."""
    try:
        items = "".join(
            f'<div class="item"><span class="label">{html.escape(k)}:</span> {html.escape(str(v))}</div>'
            for k, v in summary.items()
        )
        body = EMAIL_TEMPLATE.format(
            color="#4caf50" if ok else "#d32f2f",
            title="✅ HOÀN THÀNH" if ok else "❌ LỖI XẢY RA",
            job=JOB_NAME,
            items=items,
            log=html.escape("\n".join(LOG_BUFFER)),
        )
        msg = MIMEMultipart()
        msg["From"], msg["To"] = EMAIL_FROM, EMAIL_TO
        msg["Subject"] = f"{'✅' if ok else '❌'} {JOB_NAME} - {'THÀNH CÔNG' if ok else 'LỖI'}"
        msg.attach(MIMEText(body, "html"))

        log("Đang gửi email...", "loading")
        with smtplib.SMTP("smtp.gmail.com", 587) as server:
            server.starttls()
            server.login(EMAIL_FROM, EMAIL_PASSWORD)
            server.send_message(msg)
        log("Đã gửi email", "success")
    except Exception as e:
        log(f"Không gửi được email: {e}", "warning")


# =====================================================
# TIỆN ÍCH
# =====================================================
def js_click(driver, element):
    """Click bằng JavaScript: ổn định hơn click thường với giao diện KiotViet."""
    driver.execute_script("arguments[0].click();", element)


def calculate_date_ranges(today=None):
    """
    Khoảng cần tải: từ ngày 1 đến hôm qua của tháng hiện tại.
    Riêng ngày 1 thì tải trọn tháng trước. Chia 3 đợt: 1-10, 11-20, 21-cuối tháng.
    """
    today = today or datetime.now().date()
    end_date = today - timedelta(days=1)
    start_date = end_date.replace(day=1)

    last = end_date.day
    bounds = [(1, CHUNK_DAYS), (CHUNK_DAYS + 1, 2 * CHUNK_DAYS), (2 * CHUNK_DAYS + 1, last)]
    ranges = [(s, min(e, last)) for s, e in bounds if s <= last]
    return ranges, start_date, end_date


def open_delivery_page(driver, wait):
    """Rê chuột vào menu Đơn hàng rồi mở trang Vận đơn."""
    menu = wait.until(EC.visibility_of_element_located((By.XPATH, "//span[normalize-space()='Đơn hàng']")))
    driver.execute_script("arguments[0].dispatchEvent(new MouseEvent('mouseenter', {bubbles: true}));", menu)
    time.sleep(0.5)
    link = wait.until(EC.presence_of_element_located((By.XPATH, "//a[@href='#/OrderDelivery/']")))
    js_click(driver, link)


# =====================================================
# CHROME
# =====================================================
def create_driver():
    options = webdriver.ChromeOptions()
    options.add_argument("--start-maximized")
    options.add_argument("--disable-notifications")
    options.add_experimental_option("prefs", {
        "profile.default_content_setting_values.automatic_downloads": 1,
        "download.default_directory": DOWNLOAD_DIR,
        "download.prompt_for_download": False,
        "safebrowsing.enabled": True,
    })
    return webdriver.Chrome(service=Service(ChromeDriverManager().install()), options=options)


# =====================================================
# 1. ĐĂNG NHẬP
# =====================================================
def login_kiotviet(driver):
    log(f"Truy cập: {LOGIN_URL}")
    driver.get(LOGIN_URL)
    wait = WebDriverWait(driver, 30)

    wait.until(EC.presence_of_element_located((By.XPATH, "//input[@placeholder='Tên đăng nhập']"))).send_keys(USERNAME)
    driver.find_element(By.XPATH, "//input[@placeholder='Mật khẩu']").send_keys(PASSWORD)
    log("Đã nhập tài khoản", "success")

    remember = wait.until(EC.presence_of_element_located((By.XPATH, "//label[contains(., 'Duy trì đăng nhập')]//input")))
    if remember.is_selected():
        js_click(driver, remember)
        log('Bỏ tick "Duy trì đăng nhập"', "success")

    js_click(driver, wait.until(EC.presence_of_element_located((By.XPATH, "//button[contains(., 'Quản lý')]"))))
    wait.until(EC.visibility_of_element_located((By.XPATH, "//span[normalize-space()='Đơn hàng']")))
    log("Đăng nhập thành công", "success")


# =====================================================
# 2. CHỌN CỘT HIỂN THỊ
# =====================================================
def setup_delivery_columns(driver):
    """Mở popup ẩn/hiện cột, bỏ chọn tất cả rồi chỉ chọn REQUIRED_COLUMNS."""
    wait = WebDriverWait(driver, 60)

    for attempt in (1, 2):                                   # mở popup, thử lại 1 lần nếu lỗi
        try:
            log(f"Mở trang Vận đơn và popup cột (lần {attempt})", "step")
            open_delivery_page(driver, wait)
            time.sleep(5)                                    # chờ lưới vận đơn tải xong
            toggle = wait.until(EC.presence_of_element_located(
                (By.XPATH, "//div[@id='orderDeliveryDetailColumnSelection']//span[@class='k-link']")))
            driver.execute_script("arguments[0].scrollIntoView({block: 'center'});", toggle)
            js_click(driver, toggle)
            wait.until(EC.visibility_of_element_located((By.CSS_SELECTOR, "ul.k-popup.k-menu-group")))
            log("Popup đã mở", "success")
            break
        except Exception as e:
            if attempt == 2:
                raise RuntimeError(f"Không mở được popup chọn cột: {e}")
            log(f"Lần {attempt} thất bại, thử lại sau 2 giây", "warning")
            time.sleep(2)

    checkboxes = driver.find_elements(By.CSS_SELECTOR, "input.check_row.kv-form-check-input[type='checkbox']")
    if not checkboxes:
        raise RuntimeError("Không tìm thấy checkbox nào trong popup")

    unchecked = 0
    for cb in checkboxes:
        try:
            driver.execute_script("arguments[0].scrollIntoView({block: 'nearest'});", cb)
            if cb.is_selected():
                js_click(driver, cb)
                time.sleep(0.02)
                unchecked += 1
        except Exception:
            pass
    log(f"Bỏ chọn {unchecked}/{len(checkboxes)} cột", "success")

    selected, not_found = 0, []
    for col in REQUIRED_COLUMNS:
        try:
            cb = driver.find_element(By.XPATH,
                "//input[@class='check_row kv-form-check-input' and @type='checkbox']"
                f"[following-sibling::span[@class='kv-form-check-text' and normalize-space()='{col}']]")
            driver.execute_script("arguments[0].scrollIntoView({block: 'nearest'});", cb)
            if not cb.is_selected():
                js_click(driver, cb)
                time.sleep(0.08)
            selected += 1
        except Exception:
            not_found.append(col)

    if len(not_found) > MAX_MISSING_COLUMNS:
        raise RuntimeError(f"Chỉ chọn được {selected}/{len(REQUIRED_COLUMNS)} cột. Không tìm thấy: {', '.join(not_found)}")
    if not_found:
        log(f"Không tìm thấy cột: {', '.join(not_found)}", "warning")
    log(f"Đã chọn {selected}/{len(REQUIRED_COLUMNS)} cột", "success")

    driver.find_element(By.TAG_NAME, "body").send_keys(Keys.ESCAPE)
    time.sleep(1)


# =====================================================
# 3. XUẤT MỘT ĐỢT
# =====================================================
ROW_XPATHS = [
    "//div[contains(@class, 'k-grid')]//tbody//tr",
    "//table[@role='grid']//tbody//tr",
    "//table//tbody//tr",
]
CODE_PATTERN = re.compile(r"\b(HD\d{5,}|SPXVN\w+)")   # mã hóa đơn / mã vận đơn


def wait_for_grid_data(driver):
    """Chờ lưới dữ liệu có dòng thật (dòng có nội dung hoặc có mã hóa đơn / vận đơn)."""
    start, last_log = time.time(), 0
    while time.time() - start < DATA_WAIT_SECONDS:
        elapsed = int(time.time() - start)
        for xp in ROW_XPATHS:
            try:
                rows = [r for r in driver.find_elements(By.XPATH, xp) if len(r.text.strip()) > 10]
                if rows:
                    log(f"Dữ liệu đã hiện: {len(rows)} dòng sau {elapsed}s", "success")
                    return True
            except Exception:
                pass
        try:
            if CODE_PATTERN.search(driver.find_element(By.TAG_NAME, "body").text):
                log(f"Phát hiện mã hóa đơn / vận đơn sau {elapsed}s", "success")
                return True
        except Exception:
            pass
        if elapsed - last_log >= 5:
            log(f"[{elapsed:03d}s] Đang chờ dữ liệu...")
            last_log = elapsed
        time.sleep(2)
    log(f"Hết {DATA_WAIT_SECONDS}s chưa thấy dữ liệu rõ ràng, vẫn tiếp tục xuất file", "warning")
    return False


def export_range(driver, start_day, end_day, month, file_no):
    """Chọn khoảng ngày, chờ dữ liệu, bấm Xuất file và trả về tên file sẽ tải về."""
    wait = WebDriverWait(driver, 60)
    open_delivery_page(driver, wait)
    time.sleep(1)

    try:                                                    # bỏ lọc chi nhánh = lấy tất cả chi nhánh
        js_click(driver, wait.until(EC.presence_of_element_located(
            (By.XPATH, "//li[contains(@class,'k-button')]//span[contains(@class,'k-i-close')]"))))
    except Exception:
        pass

    js_click(driver, wait.until(EC.presence_of_element_located((By.XPATH, "//a[contains(@class,'checked')]"))))
    time.sleep(1)
    js_click(driver, wait.until(EC.presence_of_element_located((By.ID, "reportsortOtherLbl"))))
    time.sleep(1)
    js_click(driver, wait.until(EC.element_to_be_clickable(
        (By.XPATH, f"//div[@id='fromDate']//a[@class='k-link' and normalize-space()='{start_day}']"))))
    js_click(driver, wait.until(EC.element_to_be_clickable(
        (By.XPATH, f"//div[@id='toDate']//a[@class='k-link' and normalize-space()='{end_day}']"))))
    driver.execute_script(
        "document.querySelectorAll('a.kv-btn-primary').forEach(e => { if (e.innerText.trim() === 'Áp dụng') e.click(); })")
    log(f"Khoảng ngày {start_day:02d}/{month:02d} → {end_day:02d}/{month:02d}", "success")

    try:
        wait.until(EC.invisibility_of_element_located((By.CSS_SELECTOR, ".k-loading-mask")))
    except Exception:
        time.sleep(3)
    wait_for_grid_data(driver)

    js_click(driver, wait.until(EC.presence_of_element_located(
        (By.XPATH, "//a[contains(@class, 'kv2BtnExport') and @title='Xuất file']"))))
    log("Đã bấm Xuất file, chờ KiotViet tạo file...", "loading")

    dl_link = WebDriverWait(driver, EXPORT_WAIT_SECONDS).until(EC.element_to_be_clickable(
        (By.XPATH, "//div[@id='importExportContent']//a[contains(@class,'noti_link') "
                   "and contains(text(),'Nhấn vào đây để tải xuống')]")))
    filename = dl_link.get_attribute("href").split("/")[-1]
    js_click(driver, dl_link)
    log(f"File #{file_no}: {filename}", "success")
    return filename


# =====================================================
# 4. CHỜ TẢI XONG
# =====================================================
def is_file_stable(path, stable_seconds=3):
    """File coi là tải xong khi dung lượng > 0 và không đổi trong stable_seconds giây."""
    last, stable = -1, 0
    while stable < stable_seconds:
        if not os.path.exists(path) or os.path.getsize(path) == 0:
            return False
        size = os.path.getsize(path)
        stable = stable + 1 if size == last else 0
        last = size
        time.sleep(1)
    return True


def wait_for_files(directory, filenames):
    start, done = time.time(), {}
    while time.time() - start < DOWNLOAD_WAIT_SECONDS:
        for fn in filenames:
            fp = os.path.join(directory, fn)
            if fn not in done and os.path.exists(fp) and is_file_stable(fp):
                done[fn] = fp
                log(f"Tải xong {fn} ({os.path.getsize(fp) / 1024 / 1024:.1f} MB) - {len(done)}/{len(filenames)}", "success")
        if len(done) == len(filenames):
            return [done[f] for f in filenames]                # giữ đúng thứ tự ngày
        if glob.glob(os.path.join(directory, "*.crdownload")) and int(time.time() - start) % 10 == 0:
            log("Chrome đang tải...", "loading")
        time.sleep(1)
    missing = [f for f in filenames if f not in done]
    raise TimeoutError(f"Hết {DOWNLOAD_WAIT_SECONDS}s, còn thiếu: {missing}")


# =====================================================
# 5. GỘP FILE
# =====================================================
def merge_excel_files(files, output):
    """Lấy file đầu làm mẫu, nối dữ liệu các file sau, dựng lại Excel Table để giữ định dạng và bộ lọc."""
    wb = load_workbook(files[0])
    ws = wb.active
    if not ws.tables:
        raise RuntimeError("File đầu tiên không có Excel Table")

    table = list(ws.tables.values())[0]
    name, style = table.name, table.tableStyleInfo.name
    ws.tables.pop(name)

    added = 0
    for f in files[1:]:
        src = load_workbook(f, data_only=True).active
        for row in src.iter_rows(min_row=2, values_only=True):   # bỏ dòng tiêu đề
            ws.append(row)
            added += 1

    new_table = Table(displayName=name, ref=f"A1:{get_column_letter(ws.max_column)}{ws.max_row}")
    new_table.tableStyleInfo = TableStyleInfo(name=style, showRowStripes=True, showColumnStripes=False)
    ws.add_table(new_table)
    wb.save(output)

    log(f"Gộp {len(files)} file: {ws.max_row - 1:,} dòng × {ws.max_column} cột (thêm {added:,} dòng)", "success")
    return ws.max_row - 1, ws.max_column


def move_file(src, target_dir, period):
    """Đặt tên theo tháng của dữ liệu và ghi đè bản cũ của cùng tháng."""
    os.makedirs(target_dir, exist_ok=True)
    dest = os.path.join(target_dir, f"{FILE_PREFIX}_{period:%Y%m}.xlsx")
    shutil.move(src, dest + ".tmp")
    os.replace(dest + ".tmp", dest)                         # thay file cũ trong một bước
    log(f"Đã lưu: {dest} ({os.path.getsize(dest) / 1024 / 1024:.2f} MB)", "success")
    return dest


# =====================================================
# MAIN
# =====================================================
def main():
    driver, started = None, datetime.now()
    try:
        check_config()
        log(f"🚀 {JOB_NAME}", "header")
        ranges, start_date, end_date = calculate_date_ranges()
        log(f"Dữ liệu {start_date:%d/%m/%Y} - {end_date:%d/%m/%Y}, chia {len(ranges)} đợt", "section")
        for i, (s, e) in enumerate(ranges, 1):
            log(f"Đợt {i}: ngày {s:02d}-{e:02d}", "bullet")

        driver = create_driver()
        log("ĐĂNG NHẬP", "section")
        login_kiotviet(driver)
        log("CHỌN CỘT", "section")
        setup_delivery_columns(driver)

        filenames = []
        for i, (s, e) in enumerate(ranges, 1):
            log(f"XUẤT ĐỢT {i}", "section")
            filenames.append(export_range(driver, s, e, start_date.month, i))
            time.sleep(3)

        log("CHỜ TẢI FILE", "section")
        files = wait_for_files(DOWNLOAD_DIR, filenames)

        log("GỘP VÀ LƯU FILE", "section")
        temp = os.path.join(DOWNLOAD_DIR, "temp_merged.xlsx")
        rows, cols = merge_excel_files(files, temp)
        final = move_file(temp, FINAL_DIR, start_date)

        minutes, seconds = divmod(int((datetime.now() - started).total_seconds()), 60)
        log(f"HOÀN THÀNH trong {minutes} phút {seconds} giây", "result")
        send_email(True, {
            "Hoàn thành lúc": datetime.now().strftime("%d/%m/%Y %H:%M:%S"),
            "Tổng thời gian": f"{minutes} phút {seconds} giây",
            "Khoảng ngày": f"{start_date:%d/%m/%Y} - {end_date:%d/%m/%Y}",
            "Số file đã tải": len(filenames),
            "Dữ liệu": f"{rows:,} dòng × {cols} cột",
            "File cuối": final,
        })

    except Exception as e:
        log("❌ LỖI - BOT DỪNG", "header")
        log(traceback.format_exc())
        send_email(False, {
            "Thời gian lỗi": datetime.now().strftime("%d/%m/%Y %H:%M:%S"),
            "Lỗi": str(e),
            "Trạng thái": "Đã dừng, chưa ghi đè file cũ",
        })
        raise

    finally:
        if driver:
            driver.quit()


if __name__ == "__main__":
    main()
